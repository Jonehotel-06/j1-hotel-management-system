"""Batch 01 task 8 — operational report exports.

Before this batch only the occupancy report had a server-side export. The
Bookings report answered ``?export=xlsx`` with the ordinary JSON envelope
under a filename ending in ``.xlsx``: the browser saved it, Excel refused to
open it, and the staff reported "the export is broken".

These tests pin the file itself — content type, disposition, sheet contents,
currency formatting and the empty-period note — not just the status code.
"""
from __future__ import annotations

from datetime import timedelta
from decimal import Decimal
from io import BytesIO

from django.utils import timezone

from apps.bookings.models import Booking

from .base import BaseAPITestCase
from .factories import make_booking, make_guest, make_room, make_room_type, make_staff

XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _sheet_rows(payload: bytes):
    from openpyxl import load_workbook

    book = load_workbook(BytesIO(payload))
    sheet = book.active
    return [[cell for cell in row] for row in sheet.iter_rows(values_only=True)], sheet


class OperationalExportTests(BaseAPITestCase):
    """The Bookings and Occupancy endpoints must return real files."""

    def setUp(self):
        super().setUp()
        self.manager = make_staff(email="exports.manager@jone.test", role="MANAGER")
        self.guest = make_guest()
        self.room_type = make_room_type(name="Export Deluxe", price=Decimal("50000.00"))
        self.room = make_room(self.room_type, number="E-101")
        self.today = timezone.localdate()
        self.start = self.today - timedelta(days=2)
        self.end = self.today + timedelta(days=2)
        self.booking = make_booking(
            self.guest,
            self.room_type,
            rooms=[self.room],
            check_in=self.today,
            check_out=self.today + timedelta(days=2),
            status=Booking.Status.CONFIRMED,
            amount_paid=Decimal("50000.00"),
            total=Decimal("100000.00"),
        )

    def _get(self, path, **params):
        self.auth(self.manager)
        query = {"start_date": self.start.isoformat(), "end_date": self.end.isoformat(), **params}
        return self.client.get(path, query)

    # -- bookings ---------------------------------------------------------
    def test_bookings_xlsx_is_a_real_workbook_not_a_json_envelope(self):
        response = self._get("/api/admin/reports/bookings/", export="xlsx")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], XLSX_TYPE)
        payload = b"".join(response.streaming_content) if response.streaming else response.content
        # The defect: a JSON body saved under an .xlsx name. A real workbook is
        # a ZIP container, so it starts with the ZIP magic number.
        self.assertTrue(payload.startswith(b"PK"), "bookings export is not a workbook")
        self.assertNotIn(b'"success"', payload[:200])

        rows, sheet = _sheet_rows(payload)
        self.assertEqual(sheet.title, "Bookings")
        self.assertEqual(rows[0][0], "J-ONE HOTEL & LODGE — Bookings report")
        self.assertEqual(rows[1][1], self.start.isoformat())
        flat = [row[0] for row in rows]
        for heading in ("Total bookings", "Cancelled", "No-shows", "Status", "Room type"):
            self.assertIn(heading, flat)

    def test_bookings_xlsx_labels_contract_value_as_money_in_the_hotel_currency(self):
        response = self._get("/api/admin/reports/bookings/", export="xlsx")
        payload = response.content
        from openpyxl import load_workbook

        sheet = load_workbook(BytesIO(payload)).active
        header = None
        for row in sheet.iter_rows():
            if row[0].value == "Room type":
                header = row[0].row
                break
        self.assertIsNotNone(header, "room type table missing")
        self.assertEqual([cell.value for cell in sheet[header]][:3], ["Room type", "Bookings", "Contract value"])
        money_cell = sheet.cell(row=header + 1, column=3)
        self.assertEqual(Decimal(str(money_cell.value)), Decimal("100000.00"))
        # The figure carries its currency and still behaves as a number.
        self.assertIn("NGN", money_cell.number_format)

    def test_bookings_pdf_is_a_real_pdf(self):
        response = self._get("/api/admin/reports/bookings/", export="pdf")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/pdf")
        self.assertTrue(response.content.startswith(b"%PDF"))
        self.assertIn(
            f'filename="jone-bookings-{self.start}-to-{self.end}.pdf"',
            response["Content-Disposition"],
        )

    def test_bookings_without_an_export_parameter_still_returns_json(self):
        response = self._get("/api/admin/reports/bookings/")
        self.assertEqual(response.status_code, 200)
        self.assertIn("application/json", response["Content-Type"])
        self.assertEqual(response.json()["data"]["total_bookings"], 1)

    def test_an_unknown_export_format_is_ignored_rather_than_guessed(self):
        response = self._get("/api/admin/reports/bookings/", export="docx")
        self.assertEqual(response.status_code, 200)
        self.assertIn("application/json", response["Content-Type"])

    # -- occupancy --------------------------------------------------------
    def test_occupancy_export_keeps_working_and_gains_the_wider_label_column(self):
        response = self._get("/api/admin/reports/occupancy/", export="xlsx")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], XLSX_TYPE)
        self.assertIn(
            f'filename="jone-occupancy-{self.start}-to-{self.end}.xlsx"',
            response["Content-Disposition"],
        )
        from openpyxl import load_workbook

        sheet = load_workbook(BytesIO(response.content)).active
        # "Rooms not arrived (no-show)" was being clipped by a 16-char column.
        self.assertGreaterEqual(sheet.column_dimensions["A"].width, 30)


class EmptyPeriodExportTests(BaseAPITestCase):
    """An empty period must say so, not render a bare header row."""

    def setUp(self):
        super().setUp()
        self.manager = make_staff(email="exports.empty@jone.test", role="MANAGER")
        self.start = timezone.localdate() - timedelta(days=400)
        self.end = self.start + timedelta(days=1)

    def _get(self, path, export):
        self.auth(self.manager)
        return self.client.get(
            path,
            {"start_date": self.start.isoformat(), "end_date": self.end.isoformat(), "export": export},
        )

    def test_empty_bookings_workbook_states_that_nothing_happened(self):
        response = self._get("/api/admin/reports/bookings/", "xlsx")
        rows, _ = _sheet_rows(response.content)
        flat = [row[0] for row in rows]
        self.assertIn("No activity was recorded in this period.", flat)
        # The note appears under every table that has no rows, so a reader is
        # never left wondering whether the export simply failed.
        self.assertGreaterEqual(flat.count("No activity was recorded in this period."), 3)

    def test_empty_bookings_pdf_still_builds(self):
        response = self._get("/api/admin/reports/bookings/", "pdf")
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.content.startswith(b"%PDF"))

    def test_empty_occupancy_workbook_reports_each_night_as_zero(self):
        # Occupancy is per calendar night, so a quiet period is a run of zeros
        # rather than an absent table — the empty-state note is correctly not
        # used here. What matters is that the nights are not silently dropped.
        response = self._get("/api/admin/reports/occupancy/", "xlsx")
        rows, _ = _sheet_rows(response.content)
        header = next(i for i, row in enumerate(rows) if row[0] == "Date")
        nights = rows[header + 1:]
        self.assertEqual(
            [str(row[0]) for row in nights],
            [self.start.isoformat(), self.end.isoformat()],
        )
        for row in nights:
            self.assertEqual(row[1], 0, "a night with no guests must report 0, not blank")


class RevenueExportTests(BaseAPITestCase):
    """The legacy payments projection answers ?export= with a file too.

    No dashboard button points here on purpose — the reports page takes its
    money figures from the ledger summary, and a second "revenue" export fed
    by the payments projection would re-create the collections-versus-income
    conflation this batch removes. The endpoint is public API, though, and it
    must not hand back JSON under an .xlsx name.
    """

    def setUp(self):
        super().setUp()
        self.manager = make_staff(email="exports.revenue@jone.test", role="MANAGER")
        self.start = timezone.localdate() - timedelta(days=3)
        self.end = timezone.localdate()

    def _get(self, export):
        self.auth(self.manager)
        return self.client.get(
            "/api/admin/reports/revenue/",
            {"start_date": self.start.isoformat(), "end_date": self.end.isoformat(), "export": export},
        )

    def test_revenue_xlsx_is_a_workbook_and_names_the_period(self):
        response = self._get("xlsx")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], XLSX_TYPE)
        self.assertTrue(response.content.startswith(b"PK"))
        self.assertIn(
            f'filename="jone-payments-{self.start}-to-{self.end}.xlsx"',
            response["Content-Disposition"],
        )

    def test_revenue_workbook_says_it_is_collections_not_income(self):
        rows, _ = _sheet_rows(self._get("xlsx").content)
        note = " ".join(str(cell) for row in rows for cell in row if cell)
        self.assertIn("collections report, not an income statement", note)
        flat = [row[0] for row in rows]
        self.assertIn("Total received", flat)
        self.assertIn("Refunds paid out", flat)
        self.assertIn("Net received", flat)

    def test_revenue_pdf_is_a_real_pdf(self):
        response = self._get("pdf")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/pdf")
        self.assertTrue(response.content.startswith(b"%PDF"))
